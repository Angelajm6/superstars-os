"""#355 — company vehicles as a fleet: mileage, service, E-ZPass tolls, car bills.

The goal (Angela): for every company car, know how many miles it has driven, when it was
last serviced and when the next service is due, which tolls its E-ZPass tag is paying,
and every bill it costs — without busywork.

  vehicle            + ezpass_tag, ezpass_account, service_every_miles, service_every_months
  vehicle_odometer   id, vehicle_id, reading_date, odometer, source (manual|service|bill),
                     source_id, recorded_by, created_at
  vehicle_service    id, vehicle_id, service_date, odometer, service_type, vendor, cost,
                     notes, created_by, created_at
  vehicle_toll       id, vehicle_id, tag, plate, txn_date, txn_time, agency, location,
                     amount, source (import|manual), dedupe_key UNIQUE, created_by, created_at
  payable            + vehicle_id   (car bills ARE payables — they show in "what needs to be
                     paid" — tagged to the car, category 'vehicle')
  payable_file       id, payable_id, file_name, file_path, file_size, mime, uploaded_by,
                     uploaded_at   (the bill's scan/PDF; read only through the gated route)

Mileage is an odometer LOG: readings come from the vehicle page and ride along with every
service and every car bill. Miles in a period = the last reading up to the end minus the last
reading before the start (else the first reading inside it).

Tolls: E-ZPass has no public API. The office downloads the account's activity statement
(CSV or Excel) and uploads it; each row is matched to a car by its E-ZPass tag (else plate),
and a dedupe key makes re-uploading the same statement harmless. Rows for tags no car
carries are reported, not stored — add the tag to the car and upload again.

Access: the company portal gate (admin / c_suite) — nothing here is field-reachable.
File paths never enter a JSON payload (CLAUDE.md PII rule).
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from flask import jsonify, request, send_file

import db_layer
import ssc_paths
from auth import current_user, requires_section

log = logging.getLogger("ssc.fleet")

SERVICE_TYPES = ("oil_change", "tires", "brakes", "inspection", "repair", "maintenance", "other")
BILL_KINDS = ("repair", "service", "insurance", "registration", "fuel", "parking", "toll_violation", "other")
DEFAULT_EVERY_MILES = 5000
DEFAULT_EVERY_MONTHS = 6
SOON_MILES, SOON_DAYS = 500, 30
MAX_FILE_BYTES = 25 * 1024 * 1024
_EXT_MIME = {".pdf": "application/pdf", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
             ".heic": "image/heic", ".webp": "image/webp"}
_FILE_BASE = ssc_paths.under_root("data_room", "payables")


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _one(cur) -> Optional[dict]:
    r = cur.fetchone()
    return dict(r) if r else None


def _rows(cur) -> list:
    return [dict(r) for r in cur.fetchall()]


def _uid():
    return (current_user() or {}).get("id")


def _has_column(conn, table, column) -> bool:
    if db_layer.is_postgres():
        return bool(conn.execute("SELECT 1 FROM information_schema.columns WHERE table_schema='public' "
                                 "AND table_name = ? AND column_name = ?", (table, column)).fetchone())
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name = ?", (table,)).fetchone():
        return False
    return any(r[1] == column for r in conn.execute(f"PRAGMA table_info({table})").fetchall())


def _has_table(conn, table) -> bool:
    if db_layer.is_postgres():
        return bool(conn.execute("SELECT 1 FROM information_schema.tables WHERE table_schema='public' AND table_name = ?",
                                 (table,)).fetchone())
    return bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name = ?", (table,)).fetchone())


def _iso(v) -> Optional[str]:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%Y/%m/%d", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %I:%M:%S %p",
                "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M %p", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _int(v) -> Optional[int]:
    if v in (None, ""):
        return None
    try:
        return int(round(float(str(v).replace(",", "").strip())))
    except ValueError:
        return None


def _money(v) -> Optional[float]:
    if v in (None, ""):
        return None
    s = str(v).strip().replace("$", "").replace(",", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").strip()
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


# ---------------------------------------------------------------- schema
def ensure_schema(conn) -> None:
    pg = db_layer.is_postgres()
    pk = "id INTEGER GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY" if pg else "id INTEGER PRIMARY KEY AUTOINCREMENT"
    if _has_table(conn, "vehicle"):
        for col, ddl in (("ezpass_tag", "TEXT"), ("ezpass_account", "TEXT"),
                         ("service_every_miles", "INTEGER"), ("service_every_months", "INTEGER")):
            if not _has_column(conn, "vehicle", col):
                conn.execute(f"ALTER TABLE vehicle ADD COLUMN {col} {ddl}")
    if _has_table(conn, "payable") and not _has_column(conn, "payable", "vehicle_id"):
        conn.execute("ALTER TABLE payable ADD COLUMN vehicle_id INTEGER")
    conn.execute(f"CREATE TABLE IF NOT EXISTS vehicle_odometer ({pk}, vehicle_id INTEGER NOT NULL, reading_date TEXT NOT NULL, "
                 f"odometer INTEGER NOT NULL, source TEXT NOT NULL DEFAULT 'manual', source_id INTEGER, "
                 f"recorded_by INTEGER, created_at TEXT NOT NULL)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_vehicle_odometer ON vehicle_odometer(vehicle_id, reading_date)")
    conn.execute(f"CREATE TABLE IF NOT EXISTS vehicle_service ({pk}, vehicle_id INTEGER NOT NULL, service_date TEXT NOT NULL, "
                 f"odometer INTEGER, service_type TEXT NOT NULL, vendor TEXT, cost REAL, notes TEXT, "
                 f"created_by INTEGER, created_at TEXT NOT NULL)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_vehicle_service ON vehicle_service(vehicle_id, service_date)")
    conn.execute(f"CREATE TABLE IF NOT EXISTS vehicle_toll ({pk}, vehicle_id INTEGER, tag TEXT, plate TEXT, "
                 f"txn_date TEXT NOT NULL, txn_time TEXT, agency TEXT, location TEXT, amount REAL NOT NULL, "
                 f"source TEXT NOT NULL DEFAULT 'import', dedupe_key TEXT, created_by INTEGER, created_at TEXT NOT NULL)")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_vehicle_toll_dedupe ON vehicle_toll(dedupe_key)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_vehicle_toll ON vehicle_toll(vehicle_id, txn_date)")
    conn.execute(f"CREATE TABLE IF NOT EXISTS payable_file ({pk}, payable_id INTEGER NOT NULL, file_name TEXT NOT NULL, "
                 f"file_path TEXT NOT NULL, file_size INTEGER, mime TEXT, uploaded_by INTEGER, uploaded_at TEXT NOT NULL)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_payable_file ON payable_file(payable_id)")
    conn.commit()


# ---------------------------------------------------------------- mileage
def _vehicle(conn, vid: int) -> Optional[dict]:
    return _one(conn.execute("SELECT * FROM vehicle WHERE id = ?", (vid,)))


def add_reading(conn, vid: int, reading_date, odometer, *, source="manual", source_id=None, by=None) -> dict:
    d, odo = _iso(reading_date) or date.today().isoformat(), _int(odometer)
    if odo is None or odo < 0:
        raise ValueError("enter the odometer reading in miles")
    if odo > 2_000_000:
        raise ValueError("that odometer reading looks wrong")
    prev = _one(conn.execute("SELECT odometer, reading_date FROM vehicle_odometer WHERE vehicle_id = ? AND reading_date <= ? "
                             "ORDER BY reading_date DESC, id DESC LIMIT 1", (vid, d)))
    if prev and odo < prev["odometer"]:
        raise ValueError(f"lower than the {prev['odometer']:,} mi recorded on {prev['reading_date']} — check the number")
    conn.execute("INSERT INTO vehicle_odometer (vehicle_id, reading_date, odometer, source, source_id, recorded_by, created_at) "
                 "VALUES (?,?,?,?,?,?,?)", (vid, d, odo, source, source_id, by, _now()))
    return {"reading_date": d, "odometer": odo}


def _reading_at(conn, vid: int, on: str, *, before: bool = False) -> Optional[int]:
    op = "<" if before else "<="
    r = _one(conn.execute(f"SELECT odometer FROM vehicle_odometer WHERE vehicle_id = ? AND reading_date {op} ? "
                          f"ORDER BY reading_date DESC, id DESC LIMIT 1", (vid, on)))
    return r["odometer"] if r else None


def miles_between(conn, vid: int, start: str, end: str) -> Optional[int]:
    last = _reading_at(conn, vid, end)
    if last is None:
        return None
    base = _reading_at(conn, vid, start, before=True)
    if base is None:
        r = _one(conn.execute("SELECT odometer FROM vehicle_odometer WHERE vehicle_id = ? AND reading_date >= ? AND "
                              "reading_date <= ? ORDER BY reading_date, id LIMIT 1", (vid, start, end)))
        base = r["odometer"] if r else None
    return max(0, last - base) if base is not None else None


# ---------------------------------------------------------------- summary
def _add_months(d: date, months: int) -> date:
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    last = [31, 29 if (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0)) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(y, m, min(d.day, last))


def _per_month(conn, vid: int, today: date, miles: Optional[int]) -> Optional[int]:
    """Average miles a month over the span the readings actually cover (up to 12 months) —
    a car logged for 4 months is not averaged over 12."""
    if not miles:
        return None
    lo = (today - timedelta(days=365)).isoformat()
    r = _one(conn.execute("SELECT MIN(reading_date) AS a, MAX(reading_date) AS b FROM vehicle_odometer "
                          "WHERE vehicle_id = ? AND reading_date <= ?", (vid, today.isoformat())))
    if not r or not r["a"]:
        return None
    start = max(date.fromisoformat(r["a"]), date.fromisoformat(lo))
    days = (date.fromisoformat(r["b"]) - start).days
    return round(miles / max(1.0, days / 30.44))


def summary(conn, v: dict, today: Optional[date] = None) -> dict:
    today = today or date.today()
    vid = v["id"]
    latest = _one(conn.execute("SELECT odometer, reading_date FROM vehicle_odometer WHERE vehicle_id = ? "
                               "ORDER BY reading_date DESC, id DESC LIMIT 1", (vid,)))
    t = today.isoformat()
    m0, y0 = today.replace(day=1).isoformat(), today.replace(month=1, day=1).isoformat()
    last12 = miles_between(conn, vid, (today - timedelta(days=365)).isoformat(), t)
    every_mi = v.get("service_every_miles") or DEFAULT_EVERY_MILES
    every_mo = v.get("service_every_months") or DEFAULT_EVERY_MONTHS
    svc = _one(conn.execute("SELECT service_date, odometer, service_type FROM vehicle_service WHERE vehicle_id = ? "
                            "ORDER BY service_date DESC, id DESC LIMIT 1", (vid,)))
    due_date = due_miles = None
    state = "unknown"
    if svc:
        due_date = _add_months(date.fromisoformat(svc["service_date"]), every_mo).isoformat()
        if svc.get("odometer") is not None:
            due_miles = svc["odometer"] + every_mi
        odo = latest["odometer"] if latest else None
        if due_date <= t or (due_miles is not None and odo is not None and odo >= due_miles):
            state = "overdue"
        elif (date.fromisoformat(due_date) - today).days <= SOON_DAYS or \
                (due_miles is not None and odo is not None and due_miles - odo <= SOON_MILES):
            state = "soon"
        else:
            state = "ok"
    tolls = _one(conn.execute(
        "SELECT COALESCE(SUM(CASE WHEN txn_date >= ? THEN amount ELSE 0 END), 0) AS month, "
        "COALESCE(SUM(CASE WHEN txn_date >= ? THEN amount ELSE 0 END), 0) AS ytd, "
        "COALESCE(SUM(CASE WHEN txn_date >= ? THEN 1 ELSE 0 END), 0) AS n_month, MAX(txn_date) AS last "
        "FROM vehicle_toll WHERE vehicle_id = ?", (m0, y0, m0, vid))) or {}
    bills = {"open": 0.0, "open_count": 0, "ytd": 0.0}
    if _has_column(conn, "payable", "vehicle_id"):
        b = _one(conn.execute(
            "SELECT COALESCE(SUM(CASE WHEN status = 'open' THEN amount - paid_amount ELSE 0 END), 0) AS open_amt, "
            "COALESCE(SUM(CASE WHEN status = 'open' THEN 1 ELSE 0 END), 0) AS open_n, "
            "COALESCE(SUM(CASE WHEN status <> 'void' AND issued_at >= ? THEN amount ELSE 0 END), 0) AS ytd "
            "FROM payable WHERE vehicle_id = ?", (y0, vid))) or {}
        bills = {"open": round(float(b.get("open_amt") or 0), 2), "open_count": int(b.get("open_n") or 0),
                 "ytd": round(float(b.get("ytd") or 0), 2)}
    return {
        "odometer": latest["odometer"] if latest else None, "odometer_date": latest["reading_date"] if latest else None,
        "miles_month": miles_between(conn, vid, m0, t), "miles_ytd": miles_between(conn, vid, y0, t),
        "miles_12m": last12, "miles_per_month": _per_month(conn, vid, today, last12),
        "service": {"last_date": svc["service_date"] if svc else None, "last_odometer": svc.get("odometer") if svc else None,
                    "last_type": svc["service_type"] if svc else None, "due_date": due_date, "due_miles": due_miles,
                    "every_miles": every_mi, "every_months": every_mo, "state": state},
        "tolls": {"month": round(float(tolls.get("month") or 0), 2), "ytd": round(float(tolls.get("ytd") or 0), 2),
                  "count_month": int(tolls.get("n_month") or 0), "last": tolls.get("last")},
        "bills": bills,
    }


def detail(conn, vid: int, today: Optional[date] = None) -> Optional[dict]:
    v = _vehicle(conn, vid)
    if not v:
        return None
    out = {"vehicle": {k: v.get(k) for k in ("id", "label", "plate", "ezpass_tag", "ezpass_account",
                                              "service_every_miles", "service_every_months")},
           "summary": summary(conn, v, today)}
    out["readings"] = _rows(conn.execute("SELECT id, reading_date, odometer, source FROM vehicle_odometer WHERE vehicle_id = ? "
                                         "ORDER BY reading_date DESC, id DESC LIMIT 50", (vid,)))
    out["services"] = _rows(conn.execute("SELECT id, service_date, odometer, service_type, vendor, cost, notes FROM vehicle_service "
                                         "WHERE vehicle_id = ? ORDER BY service_date DESC, id DESC LIMIT 50", (vid,)))
    out["tolls"] = _rows(conn.execute("SELECT id, txn_date, txn_time, agency, location, amount, source FROM vehicle_toll "
                                      "WHERE vehicle_id = ? ORDER BY txn_date DESC, txn_time DESC, id DESC LIMIT 100", (vid,)))
    bills = []
    if _has_column(conn, "payable", "vehicle_id"):
        bills = _rows(conn.execute("SELECT id, vendor, description, doc_number, amount, paid_amount, issued_at, due_at, paid_at, status "
                                   "FROM payable WHERE vehicle_id = ? ORDER BY issued_at DESC, id DESC LIMIT 100", (vid,)))
        files = {}
        if bills:
            marks = ",".join("?" * len(bills))
            for r in _rows(conn.execute(f"SELECT id, payable_id, file_name FROM payable_file WHERE payable_id IN ({marks}) "
                                        f"ORDER BY id", [b["id"] for b in bills])):
                files.setdefault(r["payable_id"], []).append(
                    {"id": r["id"], "file_name": r["file_name"], "file_url": f"/api/company/portal/payables/files/{r['id']}"})
        for b in bills:
            b["files"] = files.get(b["id"], [])
    out["bills"] = bills
    # miles by month, last 12 months (oldest first) — the "how much is it driven" view
    months = []
    first = today or date.today()
    for i in range(11, -1, -1):
        m = _add_months(first.replace(day=1), -i)
        end = _add_months(m, 1) - timedelta(days=1)
        months.append({"month": m.strftime("%Y-%m"), "miles": miles_between(conn, vid, m.isoformat(), end.isoformat())})
    out["miles_by_month"] = months
    return out


# ---------------------------------------------------------------- service
def add_service(conn, vid: int, p: dict, *, by=None) -> dict:
    d = _iso(p.get("service_date")) or date.today().isoformat()
    st = (p.get("service_type") or "").strip()
    if st not in SERVICE_TYPES:
        raise ValueError("pick what was done")
    odo = _int(p.get("odometer"))
    cost = _money(p.get("cost"))
    conn.execute("INSERT INTO vehicle_service (vehicle_id, service_date, odometer, service_type, vendor, cost, notes, created_by, "
                 "created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                 (vid, d, odo, st, (p.get("vendor") or "").strip()[:200] or None, cost,
                  (p.get("notes") or "").strip()[:1000] or None, by, _now()))
    sid = _one(conn.execute("SELECT MAX(id) AS m FROM vehicle_service WHERE vehicle_id = ?", (vid,)))["m"]
    if odo is not None:
        add_reading(conn, vid, d, odo, source="service", source_id=sid, by=by)
    return {"id": sid}


# ---------------------------------------------------------------- tolls
def _norm_tag(s) -> str:
    return re.sub(r"\D", "", str(s or "")).lstrip("0")


def _norm_plate(s) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def _read_table(filename: str, data: bytes) -> list:
    """Rows as dicts keyed by lower-cased header, from CSV or XLSX. The header is the first row
    that looks like one (statements often start with account info lines)."""
    ext = Path(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = [list(r) for r in wb.active.iter_rows(values_only=True)]
    elif ext in (".csv", ".txt", ""):
        text = data.decode("utf-8-sig", errors="replace")
        raw = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("upload the statement as CSV or Excel (.xlsx)")
    hdr_i = None
    for i, row in enumerate(raw[:40]):
        cells = [str(c or "").strip().lower() for c in row]
        if any("date" in c for c in cells) and any(k in c for c in cells for k in ("amount", "toll", "debit", "charge")):
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("couldn't find the column headers (a date and an amount column) in that file")
    hdr = [str(c or "").strip().lower() for c in raw[hdr_i]]
    return [dict(zip(hdr, r)) for r in raw[hdr_i + 1:] if any(str(c or "").strip() for c in r)]


def _pick(hdr: list, *want, avoid=()) -> Optional[str]:
    for w in want:
        for h in hdr:
            if w in h and not any(a in h for a in avoid):
                return h
    return None


def parse_statement(filename: str, data: bytes) -> list:
    rows = _read_table(filename, data)
    if not rows:
        return []
    hdr = list(rows[0].keys())
    c_date = _pick(hdr, "transaction date", "trans date", "trans. date", "txn date", "entry date", "exit date", "date", avoid=("post",)) \
        or _pick(hdr, "date")
    c_time = _pick(hdr, "transaction time", "trans time", "exit time", "time")
    c_tag = _pick(hdr, "tag", "transponder", "device")
    c_plate = _pick(hdr, "plate")
    c_agency = _pick(hdr, "agency", "authority")
    c_entry = _pick(hdr, "entry plaza", "entry location", "entry")
    c_exit = _pick(hdr, "exit plaza", "exit location", "exit", avoid=("time", "date"))
    c_loc = _pick(hdr, "plaza", "location", "facility", "description", avoid=("entry", "exit"))
    c_amt = _pick(hdr, "amount", "toll", "debit", "charge", avoid=("balance", "time", "date", "plaza", "agency"))
    if not c_date or not c_amt:
        raise ValueError("couldn't find the date and amount columns in that file")
    out = []
    for r in rows:
        amt = _money(r.get(c_amt))
        d = _iso(r.get(c_date))
        tag, plate = (r.get(c_tag) if c_tag else None), (r.get(c_plate) if c_plate else None)
        if amt is None or not d or amt == 0 or not (_norm_tag(tag) or _norm_plate(plate)):
            continue   # payments / replenishments / blank lines carry no tag or plate
        loc = " → ".join(x for x in (str(r.get(c_entry) or "").strip() if c_entry else "",
                                     str(r.get(c_exit) or "").strip() if c_exit else "") if x) \
            or (str(r.get(c_loc) or "").strip() if c_loc else "")
        tm = r.get(c_time) if c_time else None
        if isinstance(tm, datetime):
            tm = tm.strftime("%H:%M:%S")
        out.append({"txn_date": d, "txn_time": (str(tm).strip()[:16] if tm not in (None, "") else None),
                    "tag": _norm_tag(tag) or None, "plate": _norm_plate(plate) or None,
                    "agency": (str(r.get(c_agency) or "").strip()[:80] if c_agency else None) or None,
                    "location": loc[:200] or None, "amount": round(abs(amt), 2)})
    return out


def import_tolls(conn, filename: str, data: bytes, *, by=None) -> dict:
    parsed = parse_statement(filename, data)
    cars = _rows(conn.execute("SELECT id, label, plate, ezpass_tag FROM vehicle WHERE archived_at IS NULL"))
    by_tag = {_norm_tag(c["ezpass_tag"]): c for c in cars if _norm_tag(c.get("ezpass_tag"))}
    by_plate = {_norm_plate(c["plate"]): c for c in cars if _norm_plate(c.get("plate"))}
    seen, imported, dupes, unmatched, per_car = {}, 0, 0, {}, {}
    for t in parsed:
        car = by_tag.get(t["tag"] or "") or by_plate.get(t["plate"] or "")
        if not car:
            k = ("tag …" + t["tag"][-4:]) if t["tag"] else ("plate " + t["plate"])
            unmatched[k] = unmatched.get(k, 0) + 1
            continue
        base = "|".join(str(t[k] or "") for k in ("tag", "plate", "txn_date", "txn_time", "location", "amount"))
        seen[base] = seen.get(base, 0) + 1
        key = hashlib.sha1(f"{base}|{seen[base]}".encode()).hexdigest()
        if conn.execute("SELECT 1 FROM vehicle_toll WHERE dedupe_key = ?", (key,)).fetchone():
            dupes += 1
            continue
        conn.execute("INSERT INTO vehicle_toll (vehicle_id, tag, plate, txn_date, txn_time, agency, location, amount, source, "
                     "dedupe_key, created_by, created_at) VALUES (?,?,?,?,?,?,?,?,'import',?,?,?)",
                     (car["id"], t["tag"], t["plate"], t["txn_date"], t["txn_time"], t["agency"], t["location"], t["amount"],
                      key, by, _now()))
        imported += 1
        pc = per_car.setdefault(car["id"], {"vehicle": car["label"], "tolls": 0, "amount": 0.0})
        pc["tolls"] += 1
        pc["amount"] = round(pc["amount"] + t["amount"], 2)
    conn.commit()
    return {"rows": len(parsed), "imported": imported, "duplicates": dupes,
            "unmatched": [{"who": k, "tolls": n} for k, n in sorted(unmatched.items())], "by_vehicle": list(per_car.values())}


def add_toll(conn, vid: int, p: dict, *, by=None) -> dict:
    d, amt = _iso(p.get("txn_date")), _money(p.get("amount"))
    if not d or amt is None or amt <= 0:
        raise ValueError("a toll needs a date and an amount")
    conn.execute("INSERT INTO vehicle_toll (vehicle_id, txn_date, txn_time, agency, location, amount, source, dedupe_key, "
                 "created_by, created_at) VALUES (?,?,?,?,?,?,'manual',?,?,?)",
                 (vid, d, (p.get("txn_time") or "").strip()[:16] or None, (p.get("agency") or "").strip()[:80] or None,
                  (p.get("location") or "").strip()[:200] or None, round(amt, 2), "manual:" + uuid.uuid4().hex, by, _now()))
    conn.commit()
    return {"ok": True}


# ---------------------------------------------------------------- car bills (= payables)
def add_bill(conn, vid: int, form: dict, fs=None, *, by=None) -> dict:
    import payables
    v = _vehicle(conn, vid)
    if not v:
        raise LookupError("vehicle not found")
    kind = (form.get("kind") or "other").strip()
    if kind not in BILL_KINDS:
        kind = "other"
    desc = (form.get("description") or "").strip()
    label = kind.replace("_", " ").capitalize()
    fields = {"vendor": form.get("vendor"), "amount": form.get("amount"), "issued_at": form.get("issued_at") or None,
              "due_at": form.get("due_at") or None, "doc_number": form.get("doc_number") or None, "category": "vehicle",
              "description": f"{label} · {v['label']}" + (f" — {desc}" if desc else "")}
    ext = None
    if fs is not None and fs.filename:
        ext = Path(fs.filename).suffix.lower()
        if ext not in _EXT_MIME:
            raise ValueError("attach the bill as PDF, JPG, PNG, HEIC or WEBP")
    bill = payables.create(conn, fields, source="bill", created_by=by)
    conn.execute("UPDATE payable SET vehicle_id = ? WHERE id = ?", (vid, bill["id"]))
    if str(form.get("paid") or "").lower() in ("1", "true", "on", "yes"):
        payables.pay(conn, bill["id"], paid_on=form.get("paid_on") or form.get("issued_at") or None,
                     method=form.get("method") or None)
    if form.get("odometer"):
        add_reading(conn, vid, form.get("issued_at") or date.today().isoformat(), form.get("odometer"),
                    source="bill", source_id=bill["id"], by=by)
    if ext:
        _save_file(conn, bill["id"], fs, ext, by=by)
    conn.commit()
    return payables.get(conn, bill["id"])


def _save_file(conn, payable_id: int, fs, ext: str, *, by=None) -> None:
    base = _FILE_BASE.resolve()
    ddir = _FILE_BASE / str(int(payable_id))
    if not ddir.resolve().is_relative_to(base):
        raise ValueError("invalid path")
    ddir.mkdir(parents=True, exist_ok=True)
    fpath = ddir / (uuid.uuid4().hex + ext)
    fs.save(str(fpath))
    size = fpath.stat().st_size
    if size == 0 or size > MAX_FILE_BYTES:
        fpath.unlink(missing_ok=True)
        raise ValueError("the file is empty" if size == 0 else "the file is larger than 25 MB")
    conn.execute("INSERT INTO payable_file (payable_id, file_name, file_path, file_size, mime, uploaded_by, uploaded_at) "
                 "VALUES (?,?,?,?,?,?,?)", (payable_id, Path(fs.filename).name[:200], ssc_paths.store_rel(fpath), size,
                                            _EXT_MIME[ext], by, _now()))


def resolve_file(conn, file_id: int) -> Optional[tuple]:
    r = _one(conn.execute("SELECT file_path, mime, file_name FROM payable_file WHERE id = ?", (file_id,)))
    if not r:
        return None
    p = ssc_paths.resolve_data_path(r["file_path"])
    if not (p.resolve().is_relative_to(_FILE_BASE.resolve()) and p.exists()):
        return None
    return p, r["mime"], r["file_name"]


def update_fleet_fields(conn, vid: int, p: dict) -> None:
    sets, vals = [], []
    for k in ("ezpass_tag", "ezpass_account"):
        if k in p:
            sets.append(f"{k} = ?"); vals.append((str(p.get(k) or "").strip()[:40]) or None)
    for k, lo, hi in (("service_every_miles", 500, 50000), ("service_every_months", 1, 36)):
        if k in p:
            n = _int(p.get(k))
            if n is not None and not (lo <= n <= hi):
                raise ValueError(f"{k.replace('_', ' ')} must be between {lo} and {hi}")
            sets.append(f"{k} = ?"); vals.append(n)
    if sets:
        conn.execute(f"UPDATE vehicle SET {', '.join(sets)}, updated_at = ? WHERE id = ?", vals + [_now(), vid])


# ---------------------------------------------------------------- routes
def register(app) -> None:
    conn = db_layer.connect()
    try:
        ensure_schema(conn)
        db_layer._PKMAP_CACHE.clear()
    except Exception as e:
        log.warning("#355 fleet schema ensure skipped: %s", e)
    finally:
        conn.close()
    import company_portal
    gate = requires_section(company_portal.SECTION)

    def run(fn, code=200):
        conn = db_layer.connect()
        try:
            out = fn(conn)
            conn.commit()
            return jsonify({"data": out}), code
        except LookupError as e:
            conn.rollback()
            return jsonify({"error": str(e)}), 404
        except ValueError as e:
            conn.rollback()
            return jsonify({"error": str(e)}), 400
        finally:
            conn.close()

    def need(conn, vid):
        if not _vehicle(conn, vid):
            raise LookupError("vehicle not found")

    base = "/api/company/portal/vehicles/<int:vid>"

    def _detail(vid):
        def f(c):
            d = detail(c, vid)
            if d is None:
                raise LookupError("vehicle not found")
            return d
        return run(f)

    def _fleet_fields(vid):
        p = request.get_json(silent=True) or {}
        return run(lambda c: (need(c, vid), update_fleet_fields(c, vid, p), detail(c, vid))[2])

    def _reading(vid):
        p = request.get_json(silent=True) or {}
        return run(lambda c: (need(c, vid), add_reading(c, vid, p.get("reading_date"), p.get("odometer"), by=_uid()),
                              detail(c, vid))[2], 201)

    def _service(vid):
        p = request.get_json(silent=True) or {}
        return run(lambda c: (need(c, vid), add_service(c, vid, p, by=_uid()), detail(c, vid))[2], 201)

    def _toll(vid):
        p = request.get_json(silent=True) or {}
        return run(lambda c: (need(c, vid), add_toll(c, vid, p, by=_uid()), detail(c, vid))[2], 201)

    def _bill(vid):
        form = request.form.to_dict()
        fs = request.files.get("file")
        return run(lambda c: (add_bill(c, vid, form, fs, by=_uid()), detail(c, vid))[1], 201)

    def _delete(kind, vid, row_id):
        table = {"readings": "vehicle_odometer", "services": "vehicle_service", "tolls": "vehicle_toll"}[kind]
        return run(lambda c: (c.execute(f"DELETE FROM {table} WHERE id = ? AND vehicle_id = ?", (row_id, vid)), detail(c, vid))[1])

    def _import():
        fs = request.files.get("file")
        if not fs or not fs.filename:
            return jsonify({"error": "choose the E-ZPass statement file"}), 400
        data = fs.read(10 * 1024 * 1024 + 1)
        if len(data) > 10 * 1024 * 1024:
            return jsonify({"error": "that file is larger than 10 MB"}), 400
        return run(lambda c: import_tolls(c, fs.filename, data, by=_uid()))

    def _file(file_id):
        conn = db_layer.connect()
        try:
            hit = resolve_file(conn, file_id)
        finally:
            conn.close()
        if not hit:
            return jsonify({"error": "file not found"}), 404
        p, mime, name = hit
        return send_file(str(p), mimetype=mime, download_name=name, as_attachment=False)

    app.add_url_rule(f"{base}/fleet", "fleet_detail", gate(_detail), methods=["GET"])
    app.add_url_rule(f"{base}/fleet", "fleet_fields", gate(_fleet_fields), methods=["PATCH"])
    app.add_url_rule(f"{base}/readings", "fleet_reading_add", gate(_reading), methods=["POST"])
    app.add_url_rule(f"{base}/services", "fleet_service_add", gate(_service), methods=["POST"])
    app.add_url_rule(f"{base}/tolls", "fleet_toll_add", gate(_toll), methods=["POST"])
    app.add_url_rule(f"{base}/bills", "fleet_bill_add", gate(_bill), methods=["POST"])
    app.add_url_rule(f"{base}/<kind>/<int:row_id>", "fleet_row_delete",
                     gate(lambda vid, kind, row_id: _delete(kind, vid, row_id) if kind in ("readings", "services", "tolls")
                          else (jsonify({"error": "not found"}), 404)), methods=["DELETE"])
    app.add_url_rule("/api/company/portal/vehicles/tolls/import", "fleet_toll_import", gate(_import), methods=["POST"])
    app.add_url_rule("/api/company/portal/payables/files/<int:file_id>", "payable_file", gate(_file), methods=["GET"])
