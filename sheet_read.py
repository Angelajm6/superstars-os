"""Spreadsheets on the project Documents tab (2026-10-26): .xlsx / .xlsm / .xls / .csv.

ONE reader for both uses:
  grid(data, ext)  -> {"sheets": [{name, rows, n_rows, n_cols, truncated, hidden}]} for the in-app preview
                      (cells are numbers or display strings; dates MM/DD/YYYY, percents "12.5%")
  text(data, ext)  -> the workbook as plain text for Claude's read: one "## Sheet: <name>" block per
                      sheet, cells tab-separated (the same shape claude_files sends to the Claude chat)

Values are the ones Excel last saved (formulas are never evaluated, macros never run). Nothing is
written to disk. .xls (Excel 97-2003) is read with xlrd 2.x, which reads ONLY .xls; .xlsx / .xlsm
with openpyxl (read-only, streamed).
"""
import csv
import datetime as dt
import io
import re

EXTS = (".xlsx", ".xlsm", ".xls", ".csv")
PREVIEW_ROWS = 2000      # per sheet in the preview; Download always has the full file
PREVIEW_COLS = 60
TEXT_ROWS = 5000         # per sheet in Claude's read
TEXT_MAX_CHARS = 200_000  # about 50k tokens, same cap as claude_files
MAX_SHEETS = 40

_NUM = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$|^-?\d+(\.\d+)?$")


def is_sheet(filename):
    name = str(filename or "").lower()
    return any(name.endswith(e) for e in EXTS)


def _num(v):
    """A float that is a whole number shows as one (12.0 -> 12); float noise is rounded off."""
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return str(v)
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
        return float(f"{v:.10g}")
    return v


def _date(v):
    if isinstance(v, dt.datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%m/%d/%Y")
        return v.strftime("%m/%d/%Y %I:%M %p").replace(" 0", " ")
    if isinstance(v, dt.date):
        return v.strftime("%m/%d/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%I:%M %p").lstrip("0")
    return v


def _cell(v, fmt=""):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return _date(v)
    if isinstance(v, (int, float)):
        if fmt and "%" in fmt:
            return f"{_num(v * 100)}%"
        return _num(v)
    return str(v)


def _trim(rows):
    """Drop trailing empty cells on each row and trailing empty rows (Excel often reports
    a used range far past the data)."""
    out = []
    for r in rows:
        r = list(r)
        while r and r[-1] == "":
            r.pop()
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


def _xlsx_sheets(data, max_rows, max_cols):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows, more = [], False
            for i, row in enumerate(ws.iter_rows(max_col=max_cols)):
                if i >= max_rows:
                    more = True
                    break
                rows.append([_cell(c.value, getattr(c, "number_format", "") or "") for c in row])
            yield ws.title, rows, more, getattr(ws, "sheet_state", "visible") != "visible"
    finally:
        wb.close()


def _xls_sheets(data, max_rows, max_cols):
    import xlrd
    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        for idx in range(min(book.nsheets, MAX_SHEETS)):
            sh = book.sheet_by_index(idx)
            rows = []
            for r in range(min(sh.nrows, max_rows)):
                row = []
                for c in range(min(sh.ncols, max_cols)):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(_date(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                        except Exception:
                            row.append(_num(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append("TRUE" if cell.value else "FALSE")
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        row.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        row.append(_num(cell.value))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append("")
                    else:
                        row.append(str(cell.value))
                rows.append(row)
            yield sh.name, rows, sh.nrows > max_rows, getattr(sh, "visibility", 0) != 0
            book.unload_sheet(idx)
    finally:
        book.release_resources()


def _csv_sheets(data, max_rows, max_cols):
    try:
        txt = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        txt = data.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(txt[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, more = [], False
    for i, r in enumerate(csv.reader(io.StringIO(txt), dialect)):
        if i >= max_rows:
            more = True
            break
        row = []
        for v in r[:max_cols]:
            v = v.strip()
            if _NUM.match(v):
                n = float(v.replace(",", ""))
                row.append(_num(n))
            else:
                row.append(v)
        rows.append(row)
    yield "Sheet1", rows, more, False


def _sheets(data, ext, max_rows, max_cols):
    ext = ext.lower()
    if ext in (".xlsx", ".xlsm"):
        return _xlsx_sheets(data, max_rows, max_cols)
    if ext == ".xls":
        return _xls_sheets(data, max_rows, max_cols)
    if ext == ".csv":
        return _csv_sheets(data, max_rows, max_cols)
    raise ValueError(f"not a spreadsheet: {ext}")


def grid(data, ext, max_rows=PREVIEW_ROWS, max_cols=PREVIEW_COLS):
    """The workbook for the preview. Raises on a file that does not open."""
    sheets = []
    for name, rows, more, hidden in _sheets(data, ext, max_rows, max_cols):
        rows = _trim(rows)
        sheets.append({"name": str(name), "rows": rows, "n_rows": len(rows),
                       "n_cols": max((len(r) for r in rows), default=0),
                       "truncated": bool(more), "hidden": bool(hidden)})
    return {"sheets": sheets}


def text(data, ext, max_rows=TEXT_ROWS, max_cols=PREVIEW_COLS, max_chars=TEXT_MAX_CHARS):
    """The workbook as text for Claude. Raises on a file that does not open."""
    out, size = [], 0
    for name, rows, more, hidden in _sheets(data, ext, max_rows, max_cols):
        head = f"## Sheet: {name}" + (" (hidden)" if hidden else "")
        out.append(head)
        size += len(head)
        for r in _trim(rows):
            line = "\t".join(str(v) for v in r)
            if not line.strip():
                continue
            out.append(line)
            size += len(line) + 1
            if size > max_chars:
                out.append("… (the rest of the workbook is not shown)")
                return "\n".join(out)
        if more:
            out.append(f"… (more rows in this sheet not shown; first {max_rows} only)")
    return "\n".join(out)
