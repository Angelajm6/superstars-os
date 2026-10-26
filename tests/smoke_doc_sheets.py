#!/usr/bin/env python3
"""Spreadsheets on the project Documents tab (2026-10-26): behavioral guard.

Boots the real app (`import server`) on an ISOLATED database with a scratch data root.
Claude's read is the DOC_SCAN_FAKE seam for the routes, and a stand-in Anthropic client for
the content-block check, so no API key is needed. Proves:

  * sheet_read: .xlsx / .xlsm / .xls / .csv open; dates MM/DD/YYYY, percents, booleans, whole
    numbers; trailing blanks trimmed; hidden sheets flagged; big sheets capped + marked; the
    text read has one "## Sheet:" block per sheet; a broken file raises
  * Claude's read: a spreadsheet goes to the model as a TEXT document block (never as a
    picture); a broken spreadsheet is a ScanError -> the card waits as 'unread'
  * upload: .xls / .xlsm / .csv are accepted on the drop zone with their own type label + mime
  * preview: /api/documents/intake/<id>/sheet and /api/documents/<id>/sheet return the cells;
    anon 401, client 403, a PM not on the job 403; a PDF is 400; NO *_path in the payload
  * the Documents tab: every file picker takes Excel + CSV; the side panel draws a 'sheet'

Run:  venv/bin/python tests/smoke_doc_sheets.py
"""
from __future__ import annotations

import datetime as dt
import io
import json
import os
os.environ.setdefault("SSC_ALLOW_TEST_SEAMS", "1")   # test-only seams are inert without this (runtime_seams.py)
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_TMP = tempfile.mkdtemp(prefix="ssc_docsheets_")

from _isolated_db import build_isolated_sqlite as _build_isolated_sqlite, purge_users, purge  # noqa: E402

if not (os.environ.get("SSC_DB_URL") or "").strip():
    os.environ["SSC_DB_URL"] = _build_isolated_sqlite()
DATA_ROOT = Path(_TMP) / "dataroot"
DATA_ROOT.mkdir(parents=True, exist_ok=True)
os.environ["SSC_DATA_ROOT"] = str(DATA_ROOT)
os.environ.setdefault("SSC_STRICT_ROUTES", "0")
os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ.pop("DOC_SCAN_FAKE", None)

import openpyxl  # noqa: E402

import auth  # noqa: E402
import db_layer  # noqa: E402
import document_scanner as scanner  # noqa: E402
import server  # noqa: E402  (boots the whole app)
import sheet_read  # noqa: E402

TAG = "sheets26"
PW = "DocSheetsSmokePassword2026"
ROLES = ("admin", "pm", "client")
PROJ = "SHEET-SMOKE-A"
FAKE = Path(_TMP) / "scan.json"
XLS = (ROOT / "tests" / "fixtures" / "takeoff_sample.xls").read_bytes()
CELL = "Spalled lintel 4th fl"   # never in logs


def _xlsx(tag="takeoff"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Takeoff"
    ws.append(["Drop", "Elevation", "Brick repl (SF)", "Done", "Inspected", "Note"])
    ws.append([3, "North", 312.0, 0.125, dt.date(2026, 10, 14), CELL])
    ws.append([4, "East", 288, True, dt.datetime(2026, 10, 15, 9, 30), None])
    ws["D2"].number_format = "0.0%"
    ws["H40"] = None   # a used range past the data
    hid = wb.create_sheet("Rates")
    hid.sheet_state = "hidden"
    hid.append(["Item", "Unit"])
    wb.create_sheet("Empty")
    wb.properties.title = tag   # distinct bytes per upload (the drop zone de-duplicates)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _db():
    return db_layer.connect(pragma_fk=False)


def _seed():
    conn = _db()
    try:
        for role in ROLES:
            email = f"{TAG}.{role}@example.com"
            conn.execute("DELETE FROM users WHERE email = ?", (email,))
            conn.execute("INSERT INTO users (email, password_hash, role, full_name, is_active, status, "
                         "must_reset_password, is_system) VALUES (?, ?, ?, ?, 1, 'active', 0, 1)",
                         (email, auth.hash_password(PW), role, f"Smoke {role}"))
        conn.execute("DELETE FROM projects WHERE project_code = ?", (PROJ,))
        conn.execute("INSERT INTO projects (project_code, name, address, status) VALUES (?, 'Sheet Facade', "
                     "'1 Sheet Way', 'active')", (PROJ,))
        conn.commit()   # the PM is deliberately NOT assigned to this job
    finally:
        conn.close()


def _cleanup():
    conn = _db()
    try:
        doc_ids = [r["id"] for r in conn.execute("SELECT id FROM project_documents WHERE project_code = ?", (PROJ,)).fetchall()]
        for t in ("item_visibility", "item_redflag", "visibility_audit"):
            purge(conn, t, "item_id", doc_ids)
        purge(conn, "project_documents", "id", doc_ids)
        purge(conn, "project_doc_intake", "project_code", [PROJ])
        conn.execute("DELETE FROM projects WHERE project_code = ?", (PROJ,))
        uids = [r["id"] for r in conn.execute("SELECT id FROM users WHERE email LIKE ?", (f"{TAG}.%",)).fetchall()]
        conn.commit()
        purge_users(conn, uids)
        conn.commit()
    finally:
        conn.close()
    shutil.rmtree(_TMP, ignore_errors=True)


def _no_paths(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            assert not k.endswith("_path"), k
            _no_paths(v)
    elif isinstance(obj, list):
        for v in obj:
            _no_paths(v)


def _login(role):
    c = server.app.test_client()
    r = c.post("/api/auth/login", json={"email": f"{TAG}.{role}@example.com", "password": PW})
    assert r.status_code == 200, (role, r.status_code)
    return c


def _data(r):
    j = r.get_json() or {}
    return j.get("data", j)


def _drop(c, name, body):
    return c.post(f"/api/projects/{PROJ}/documents/intake",
                  data={"file": (io.BytesIO(body), name)}, content_type="multipart/form-data")


class _Trap:
    def __init__(self):
        self.lines = []

    def write(self, s):
        self.lines.append(s)

    def flush(self):
        pass


def main() -> int:
    checks = 0
    import logging
    trap = _Trap()
    h = logging.StreamHandler(trap)
    logging.getLogger().addHandler(h)
    _seed()

    # ---- sheet_read: xlsx
    g = sheet_read.grid(_xlsx(), ".xlsx")
    names = [s["name"] for s in g["sheets"]]
    assert names == ["Takeoff", "Rates", "Empty"], names
    tk = g["sheets"][0]
    assert tk["rows"][0][0] == "Drop" and tk["n_rows"] == 3 and tk["n_cols"] == 6, tk
    assert tk["rows"][1][:5] == [3, "North", 312, "12.5%", "10/14/2026"], tk["rows"][1]
    assert tk["rows"][2][2:5] == [288, "TRUE", "10/15/2026 9:30 AM"], tk["rows"][2]
    assert tk["rows"][2] == tk["rows"][2][:5], "trailing blanks trimmed"
    assert g["sheets"][1]["hidden"] and not tk["hidden"] and g["sheets"][2]["rows"] == []
    t = sheet_read.text(_xlsx(), ".xlsx")
    assert "## Sheet: Takeoff\nDrop\tElevation" in t and "## Sheet: Rates (hidden)" in t and CELL in t
    checks += 1

    # ---- xls (Excel 97-2003), csv, caps, broken files
    x = sheet_read.grid(XLS, ".xls")
    assert [s["name"] for s in x["sheets"]] == ["Takeoff", "Totals"]
    assert x["sheets"][0]["rows"][1] == [3, "North", 312, 9.2, "10/14/2026"], x["sheets"][0]["rows"][1]
    assert x["sheets"][0]["rows"][2][4] == "TRUE" and x["sheets"][1]["rows"] == [["Total CF", 20.7]]
    c = sheet_read.grid("Item;Qty;Cost\nBrick;1,284;€ 9\n".encode("cp1252"), ".csv")
    assert c["sheets"][0]["rows"] == [["Item", "Qty", "Cost"], ["Brick", 1284, "€ 9"]], c
    big = "\n".join(f"{i},{i * 2}" for i in range(50)).encode()
    capped = sheet_read.grid(big, ".csv", max_rows=10)["sheets"][0]
    assert capped["n_rows"] == 10 and capped["truncated"]
    assert "first 10 only" in sheet_read.text(big, ".csv", max_rows=10)
    assert len(sheet_read.text(big, ".csv", max_chars=40)) < 120
    for ext in (".xlsx", ".xls"):
        try:
            sheet_read.grid(b"not a workbook", ext)
            raise AssertionError(f"broken {ext} opened")
        except AssertionError:
            raise
        except Exception:
            pass
    assert sheet_read.is_sheet("A.XLS") and not sheet_read.is_sheet("a.pdf")
    checks += 1

    # ---- Claude's read: a text document block, never a picture
    spec = scanner.spec_from_bytes(_xlsx(), "Drop 3 takeoff.xlsx")
    assert spec["is_text"] and spec["title"] == "Drop 3 takeoff.xlsx" and "## Sheet: Takeoff" in spec["text"]
    assert "data_b64" not in spec
    try:
        scanner.spec_from_bytes(b"junk", "x.xls")
        raise AssertionError("broken xls became a spec")
    except scanner.ScanError:
        pass
    sent = {}

    class _Msgs:
        def create(self, **kw):
            sent.update(kw)

            class R:
                content = [type("B", (), {"type": "text", "text": json.dumps({"title": "Takeoff"})})()]
                usage = type("U", (), {"input_tokens": 1, "output_tokens": 1})()
            return R()

    import anthropic
    real = anthropic.Anthropic
    anthropic.Anthropic = lambda *a, **k: type("C", (), {"messages": _Msgs()})()
    os.environ["ANTHROPIC_API_KEY"] = "test-not-a-key"
    try:
        scanner.call_vision_model([spec, scanner.spec_from_bytes(b"%PDF-1.4", "a.pdf")], [("CONTRACTS", "C")], [])
    finally:
        anthropic.Anthropic = real
        os.environ.pop("ANTHROPIC_API_KEY", None)
    blocks = sent["messages"][0]["content"]
    assert blocks[0] == {"type": "document", "title": "Drop 3 takeoff.xlsx",
                         "source": {"type": "text", "media_type": "text/plain", "data": spec["text"]}}, blocks[0]
    assert blocks[1]["source"]["media_type"] == "application/pdf"
    assert not any(b.get("type") == "image" for b in blocks)
    assert "## Sheet: <name>" in blocks[-1]["text"]
    checks += 1

    # ---- upload + preview routes
    admin, pm, client = _login("admin"), _login("pm"), _login("client")
    anon = server.app.test_client()
    FAKE.write_text(json.dumps({"title": "Drop 3 takeoff", "category": "DRAWINGS", "requirement_key": None,
                                "type_key": "contract_other", "confidence": 0.9, "summary": "A takeoff.",
                                "key_facts": [], "fields": []}), encoding="utf-8")
    os.environ["DOC_SCAN_FAKE"] = str(FAKE)
    try:
        cards = {}
        for name, body, label, mime in (
                ("takeoff.xls", XLS, "XLS", "application/vnd.ms-excel"),
                ("takeoff.xlsm", _xlsx("xlsm"), "XLSM", "application/vnd.ms-excel.sheet.macroEnabled.12"),
                ("takeoff.csv", b"Drop,SF\n3,312\n", "CSV", "text/csv"),
                ("takeoff.xlsx", _xlsx(), "XLSX", None)):
            r = _drop(admin, name, body)
            assert r.status_code == 201, (name, r.status_code, r.get_data(as_text=True)[:200])
            card = _data(r)
            assert card["status"] == "ready" and card["doc_type"] == label, card
            if mime:
                assert card["mime"] == mime
            cards[name] = card
        assert _drop(admin, "takeoff.ods", b"x").status_code == 400

        url = f"/api/documents/intake/{cards['takeoff.xls']['id']}/sheet"
        assert anon.get(url).status_code == 401
        assert client.get(url).status_code == 403
        assert pm.get(url).status_code == 403
        r = admin.get(url)
        assert r.status_code == 200 and "no-store" in r.headers.get("Cache-Control", "")
        d = _data(r)
        _no_paths(d)
        assert d["sheets"][0]["rows"][1][1] == "North" and d["preview_rows"] == sheet_read.PREVIEW_ROWS

        pdf = _data(_drop(admin, "permit.pdf", b"%PDF-1.4 sheets %%EOF"))
        assert admin.get(f"/api/documents/intake/{pdf['id']}/sheet").status_code == 400

        r = admin.post(f"/api/documents/intake/{cards['takeoff.xlsx']['id']}/approve", json={"type_key": "contract_other"})
        assert r.status_code == 201, r.get_data(as_text=True)[:300]
        doc = _data(r)
        doc_id = doc.get("id") or doc.get("document", {}).get("id")
        surl = f"/api/documents/{doc_id}/sheet"
        assert anon.get(surl).status_code == 401
        assert client.get(surl).status_code == 403
        assert pm.get(surl).status_code == 403
        d = _data(admin.get(surl))
        _no_paths(d)
        assert [s["name"] for s in d["sheets"]] == ["Takeoff", "Rates", "Empty"]
        assert admin.get("/api/documents/99999999/sheet").status_code == 404

        # broken spreadsheet: preview says so; Claude's read leaves the card 'unread'
        os.environ.pop("DOC_SCAN_FAKE", None)
        os.environ["ANTHROPIC_API_KEY"] = "test-not-a-key"
        try:
            bad = _data(_drop(admin, "broken.xlsx", b"not a workbook"))
        finally:
            os.environ.pop("ANTHROPIC_API_KEY", None)
        assert bad["status"] == "unread", bad
        assert admin.get(f"/api/documents/intake/{bad['id']}/sheet").status_code == 422
        checks += 1
    finally:
        os.environ.pop("DOC_SCAN_FAKE", None)

    # ---- the Documents tab + side panel
    html = (ROOT / "dashboard-static.html").read_text(encoding="utf-8")
    assert html.count('accept=".pdf,.jpg,.jpeg,.png,.heic,.heif,.xlsx,.xlsm,.xls,.csv"') == 3
    assert "inp.accept='.pdf,.jpg,.jpeg,.png,.heic,.heif,.xlsx,.xlsm,.xls,.csv'" in html
    assert "sheetUrl:kind==='sheet'?String(o.url).replace(/\\/file$/,'/sheet')" in html
    js = (ROOT / "static/js/pdf_view.js").read_text(encoding="utf-8")
    assert "return 'sheet'" in js and "sheet(o.sheetUrl, host)" in js and "position:sticky" in js
    assert "esc(cellText(v))" in js and "esc(s.name)" in js
    assert "xlrd" in (ROOT / "requirements.txt").read_text(encoding="utf-8")
    checks += 1

    logging.getLogger().removeHandler(h)
    assert CELL not in "".join(trap.lines), "cell content reached the logs"
    print(f"smoke_doc_sheets: {checks} checks passed")
    return 0


if __name__ == "__main__":
    try:
        rc = main()
    finally:
        _cleanup()
    sys.exit(rc)
